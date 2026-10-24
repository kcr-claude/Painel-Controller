"""Parametros do mes (aba ANALISE) e varredura de erros de formula."""
from openpyxl.utils import get_column_letter

from app.nucleo.normaliza import chave, erro_formula, num


def ler_parametros(wb) -> dict:
    """ANALISE: 'DIA DO MES' e 'QTDE. DIAS NO MES' ficam com o valor na coluna seguinte."""
    ws = wb["ANÁLISE"] if "ANÁLISE" in wb.sheetnames else wb[wb.sheetnames[0]]
    achados = {}
    for row in ws.iter_rows(max_row=12, max_col=8, values_only=True):
        for i, v in enumerate(row[:-1]):
            if not isinstance(v, str) or v.strip().startswith("%"):
                continue  # '% DIA DO MES' e outro indicador
            k = chave(v)
            if k == "DIA DO MES":
                achados.setdefault("dia_lancado", num(row[i + 1]))
            elif k == "QTDE DIAS NO MES":
                achados.setdefault("dias_mes", num(row[i + 1]))
    return {k: int(v) if v is not None else None for k, v in achados.items()}


def varrer_erros(wb, colunas_metas=None, limite_linhas=260):
    """Lista (aba, celula, valor) com #REF!, #DIV/0!, #N/A etc.
    Em METAS so o bloco do mes do arquivo (colunas_metas = (inicio, fim), base 0)."""
    saida = []
    for aba in ("ANÁLISE", "METAS", "CUSTOS", "VENDAS"):
        if aba not in wb.sheetnames:
            continue
        ini, fim = colunas_metas if aba == "METAS" and colunas_metas else (0, 10_000)
        for r, row in enumerate(wb[aba].iter_rows(max_row=limite_linhas, values_only=True), 1):
            for c, v in enumerate(row, 1):
                if ini < c <= fim and erro_formula(v):
                    saida.append({"aba": aba, "celula": f"{get_column_letter(c)}{r}", "valor": v.strip()})
    return saida
